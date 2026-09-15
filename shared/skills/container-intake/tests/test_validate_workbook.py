import importlib.util
from pathlib import Path
import tempfile
import unittest

import openpyxl

spec = importlib.util.spec_from_file_location('validator', Path(__file__).parents[1] / 'scripts/validate_workbook.py')
v = importlib.util.module_from_spec(spec)
spec.loader.exec_module(v)


class ValidationTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / 'input.xlsx'
        self.wb = openpyxl.Workbook()
        s = self.wb.active
        s.title = 'data'
        s.append(v.HEADERS)
        s.append(['PM535', '6903070215357', '袜子', None, None, 10, 80, 2,
                  1600, 160, .1, 100, 100, 10, .2, '新疆', None, '打'])
        s.append(['DATA_END'])
        s.append(['合计', None, None, None, None, None, None, 2, 1600, 160,
                  None, None, None, None, .2])
        c = self.wb.create_sheet('config')
        c.append(['参数', '值', '说明'])
        c.append(['货柜号', 'TGBU5973558'])
        c.append(['海运费', 2000, 'USD'])
        c.append(['内陆费', 0, '厂家承担'])

    def check(self):
        self.wb.save(self.path)
        before = self.path.read_bytes()
        result = v.validate(self.path)
        self.assertEqual(before, self.path.read_bytes())
        return result

    def test_valid_and_read_only(self):
        self.assertEqual(self.check()['status'], 'PASS')

    def test_each_column_rejects_invalid_content(self):
        for col in range(1, 19):
            with self.subTest(col=col):
                cell = self.wb['data'].cell(2, col)
                original = cell.value
                cell.value = 'bad' if col in v.NUMERIC or col in (2, 18) else 123
                self.assertEqual(self.check()['status'], 'ERROR')
                cell.value = original

    def test_arithmetic_and_totals(self):
        for coordinate in ['I2', 'J2', 'H4', 'I4', 'J4', 'O4']:
            with self.subTest(cell=coordinate):
                cell = self.wb['data'][coordinate]
                original = cell.value
                cell.value += 1
                self.assertEqual(self.check()['status'], 'ERROR')
                cell.value = original

    def test_formula_cache_is_not_silently_accepted(self):
        self.wb['data']['I2'] = '=F2*J2'
        result = self.check()
        self.assertEqual(result['status'], 'PENDING')
        self.assertTrue(any(i['cell'] == 'I2' for i in result['issues']))

    def test_missing_price_and_barcode_are_pending(self):
        self.wb['data']['F2'] = None
        self.wb['data']['B2'] = None
        self.assertEqual(self.check()['status'], 'PENDING')

    def test_structure_and_config(self):
        self.wb['config']['D1'] = '说明'
        self.wb['config']['B3'] = -1
        self.wb['data']['A3'] = 'forgot marker'
        self.assertEqual(self.check()['status'], 'ERROR')

    def test_loose_stock_exceeding_standard_case_keeps_identity(self):
        s = self.wb['data']
        s['S1'] = '散件数'
        s['S2'] = 110
        s['S4'] = 110
        s['H2'] = 4
        s['I2'] = 4300
        s['J2'] = 430
        s['O2'] = .5
        s['H4'] = 4
        s['I4'] = 4300
        s['J4'] = 430
        s['O4'] = .5
        self.assertEqual(self.check()['status'], 'REVIEW')
        s['J2'] = 320
        self.assertEqual(self.check()['status'], 'ERROR')

    def test_mixed_units_require_separate_quantity_totals(self):
        s = self.wb['data']
        s['F1'] = '单价（元）'
        s.insert_rows(3)
        for c in range(1, 19):
            s.cell(3, c).value = s.cell(2, c).value
        s['A3'] = '40S-1381'
        s['C3'] = '女士内裤'
        s['R3'] = '条'
        s['H5'] = 4
        s['I5'] = 3200
        s['O5'] = .4
        s['J5'] = None
        s['A6'] = '打数合计'
        s['J6'] = 160
        s['R6'] = '打'
        s['A7'] = '条数合计'
        s['J7'] = 160
        s['R7'] = '条'
        self.assertEqual(self.check()['status'], 'PASS')
        s['J5'] = 320
        self.assertEqual(self.check()['status'], 'ERROR')

    def test_volume_difference_requires_source_judgment(self):
        self.wb['data']['O2'] = .15
        self.wb['data']['O4'] = .15
        result = self.check()
        self.assertEqual(result['status'], 'REVIEW')
        self.assertTrue(any(i['cell'] == 'O2' and i['level'] == 'review' for i in result['issues']))

    def test_pass_does_not_claim_source_correctness(self):
        # Plausible but wrong identity cannot be rejected from output alone.
        self.wb['data']['A2'] = 'WRONG-SKU'
        result = self.check()
        self.assertEqual(result['status'], 'PASS')
        self.assertEqual(result['source_review']['status'], 'NOT_PERFORMED')

    def test_layout_gap_is_review_but_misplaced_note_is_checked(self):
        c = self.wb['config']
        c.insert_cols(3)
        self.assertEqual(self.check()['status'], 'REVIEW')
        c['D3'] = 123
        self.assertEqual(self.check()['status'], 'ERROR')

    def test_mixed_loose_quantities_cannot_be_added(self):
        self.test_mixed_units_require_separate_quantity_totals()
        s = self.wb['data']
        s['J5'] = None
        s['S1'] = '散件数'
        s['S2'] = s['S3'] = 1
        s['J2'] = s['J3'] = 161
        s['I2'] = s['I3'] = 1610
        s['I5'] = 3220
        s['J6'] = s['J7'] = 161
        s['S6'] = s['S7'] = 1
        self.assertEqual(self.check()['status'], 'REVIEW')
        s['S5'] = 2
        result = self.check()
        self.assertEqual(result['status'], 'ERROR')
        self.assertTrue(any(i['cell'] == 'S5' for i in result['issues']))

    def test_standard_template_maps_units_and_loose_columns(self):
        old = self.wb['data']
        self.wb.remove(old)
        s = self.wb.create_sheet('data', 0)
        s.append(v.STANDARD_HEADERS)
        s.append(['MF401', '6903040084013', '女士袜', None, '新疆', '打', 10, 80, 4, 110, 430, 4300, .5])
        s.append(['DATA_END'])
        s.append(['合计', None, None, None, None, None, None, None, 4, 110, 430, 4300, .5])
        self.wb['config'].append(['发柜日期', '2026-09-08'])
        self.wb['config'].append(['ETA', '2026-10-18'])
        result = self.check()
        self.assertFalse(any(i['level'] in ('error', 'pending') for i in result['issues']))
        self.assertEqual(result['rows'], 1)
        s['K2'] = 320
        result = self.check()
        self.assertTrue(any(i['cell'] == 'K2' and i['level'] == 'error' for i in result['issues']))
        s['K2'] = 430
        s['F2'] = '条'
        s['C2'] = '女士内裤'
        self.assertFalse(any(i['level'] in ('error', 'pending') for i in self.check()['issues']))


if __name__ == '__main__':
    unittest.main()
