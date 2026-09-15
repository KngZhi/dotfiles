#!/usr/bin/env python3
"""Read-only column validation; no formula calculation or source inference."""
import argparse
import datetime as dt
import hashlib
import json
import math
from pathlib import Path
import re

import openpyxl

HEADERS = ['货号', '条形码', '品名', '图片', '工艺', '单价（元/打）', '装箱数',
           '件数', '总价格', '总数量', '立方', '长', '宽', '高', '总立方', '供应商',
           '备     注', '计价单位']
NUMERIC = {6, 7, 8, 9, 10, 11, 12, 13, 14, 15}
REQUIRED = {1, 3, 6, 7, 8, 10, 15, 16, 18}
FEES = {'海运费', '内陆费', '卸柜费', '清关杂费', 'IVA'}
RATES = {'CNY-CLP', 'USD-CLP', 'USD-CNY'}


def number(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)


def validate(path):
    path = Path(path)
    wb = openpyxl.load_workbook(path, data_only=False)
    cache = openpyxl.load_workbook(path, data_only=True)
    issues = []
    def issue(level, sheet, cell, message):
        issues.append(dict(level=level, sheet=sheet, cell=cell, message=message))
    def read(cell):
        if cell.data_type == 'e':
            issue('error', cell.parent.title, cell.coordinate, 'Excel 错误值')
            return None
        if cell.data_type == 'f':
            v = cache[cell.parent.title][cell.coordinate]
            if v.value is None or v.data_type == 'e':
                issue('pending', cell.parent.title, cell.coordinate, '公式无有效缓存；须重算并保存后复验')
                return None
            return v.value
        return cell.value
    def compare(s, coordinate, actual, expected, label):
        if number(actual) and not math.isclose(actual, expected, rel_tol=1e-6, abs_tol=0.01 if label == '货值' else 1e-6):
            issue('error', s, coordinate, f'{label}不符：实际 {actual}，应为 {expected}')
    if set(wb.sheetnames) != {'data', 'config'}:
        issue('error', '', '', '工作簿必须且只能包含 data、config')
    totals = {c: [] for c in (8, 9, 10, 15)}
    rows = 0
    if 'data' in wb:
        ws = wb['data']
        headers = [c.value for c in ws[1]]
        if headers != HEADERS:
            issue('error', 'data', '1', f'列顺序或标题错误；应为 {HEADERS}')
        else:
            ends = [r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == 'DATA_END']
            if len(ends) != 1:
                issue('error', 'data', 'A', '必须且只能有一个 DATA_END')
            else:
                end = ends[0]
                if end == 2:
                    issue('error', 'data', 'A2', '没有商品明细')
                for r in range(2, end):
                    rows += 1
                    vals = {}
                    for c in range(1, 19):
                        cell = ws.cell(r, c)
                        v = vals[c] = read(cell)
                        if v in (None, ''):
                            if c in REQUIRED or c == 2:
                                issue('pending', 'data', cell.coordinate, f'{HEADERS[c-1]}缺失，需来源或明确例外')
                            continue
                        if c in NUMERIC:
                            if not number(v) or v <= 0:
                                issue('error', 'data', cell.coordinate, '必须为大于 0 的数字；未知值留空')
                            elif c == 8 and int(v) != v:
                                issue('error', 'data', cell.coordinate, '件数必须是整数')
                        elif not isinstance(v, str):
                            issue('error', 'data', cell.coordinate, '必须是文本，货号和条码不得存成数字')
                        elif c == 18 and v != '打':
                            issue('error', 'data', cell.coordinate, '计价单位必须为打')
                        elif c == 2:
                            if not re.fullmatch(r'\d{13}', v) or sum(int(n) * (1 if i % 2 == 0 else 3) for i, n in enumerate(v)) % 10:
                                issue('error', 'data', cell.coordinate, 'EAN-13 格式或校验位错误')
                    for target, factors, label in [(10, (7, 8), '数量'), (9, (6, 10), '货值'), (15, (11, 8), '体积')]:
                        if all(number(vals[c]) for c in factors):
                            expected = math.prod(vals[c] for c in factors)
                            compare('data', ws.cell(r, target).coordinate, vals[target], expected, label)
                            if vals[target] in (None, '') and ws.cell(r, target).data_type != 'f':
                                issue('pending', 'data', ws.cell(r, target).coordinate, f'缺少可核对的{label}')
                    dims = [vals[c] for c in (12, 13, 14)]
                    if any(v is not None for v in dims):
                        if all(number(v) for v in dims):
                            compare('data', f'K{r}', vals[11], math.prod(dims) / 1000000, '箱体积（厘米换立方米）')
                        else:
                            issue('pending', 'data', f'L{r}:N{r}', '长宽高未完整填写')
                    for c in totals:
                        totals[c].append(vals[c])
                # Only one nonempty summary row may follow the marker.
                tail = [r for r in range(end + 1, ws.max_row + 1)
                        if any(ws.cell(r, c).value is not None for c in range(1, 19))]
                if len(tail) != 1 or ws.cell(tail[0], 1).value not in ('合计', '总计'):
                    issue('error', 'data', f'A{end+1}', 'DATA_END 后必须只有一行合计，不能隐藏商品行')
                else:
                    for c, values in totals.items():
                        cell = ws.cell(tail[0], c)
                        v = read(cell)
                        if all(number(x) for x in values):
                            if not number(v):
                                issue('pending', 'data', cell.coordinate, '合计缺少有效数字或公式缓存')
                            else:
                                compare('data', cell.coordinate, v, sum(values), '货值' if c == 9 else '合计')
    if 'config' in wb:
        ws = wb['config']
        if [c.value for c in ws[1]] != ['参数', '值', '说明']:
            issue('error', 'config', '1', '必须为 参数 / 值 / 说明 三列')
        keys = set()
        for r in range(2, ws.max_row + 1):
            if all(c.value is None for c in ws[r]):
                continue
            key, cell = ws.cell(r, 1).value, ws.cell(r, 2)
            if not isinstance(key, str) or key in keys:
                issue('error', 'config', f'A{r}', '参数名缺失或重复')
            keys.add(key)
            v = read(cell)
            if v in (None, ''):
                if key in {'货柜号', '海运费', '内陆费'}:
                    issue('pending', 'config', cell.coordinate, f'{key}待补')
            elif key in FEES | RATES:
                if not number(v) or v < 0 or (key in RATES and v == 0):
                    issue('error', 'config', cell.coordinate, '费用须非负数字，汇率须大于 0')
            elif key == '货柜号':
                if not isinstance(v, str) or not re.fullmatch(r'[A-Z]{4}\d{7}', v):
                    issue('error', 'config', cell.coordinate, '柜号须为四个大写字母加七位数字')
            elif key == 'ETA':
                try:
                    if not isinstance(v, (dt.date, dt.datetime)):
                        dt.date.fromisoformat(v)
                except (TypeError, ValueError):
                    issue('error', 'config', cell.coordinate, 'ETA 须为日期或 YYYY-MM-DD')
            else:
                issue('pending', 'config', cell.coordinate, f'参数 {key} 无自动业务规则，须人工核对')
            note = ws.cell(r, 3).value
            if note is not None and not isinstance(note, str):
                issue('error', 'config', f'C{r}', '说明必须为文本')
        for key in {'货柜号', '海运费', '内陆费'} - keys:
            issue('pending', 'config', 'A', f'缺少参数 {key}')
    result = dict(workbook=str(path.resolve()), sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
                  rows=rows, columns=HEADERS, issues=issues,
                  status='ERROR' if any(i['level'] == 'error' for i in issues) else 'PENDING' if issues else 'PASS',
                  scope='表内列契约与算术；原件归属、单价依据、条码身份、图片、压缩体积和排版须独立核对')
    wb.close()
    cache.close()
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    parser.add_argument('--report', type=Path, required=True)
    args = parser.parse_args()
    if args.report.resolve() == args.workbook.resolve():
        parser.error('报告不能覆盖输入工作簿')
    result = validate(args.workbook)
    args.report.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f"{result['status']}: {result['rows']} rows; {len(result['issues'])} issues; {args.report}")
    return {'PASS': 0, 'ERROR': 1, 'PENDING': 2}[result['status']]


if __name__ == '__main__':
    raise SystemExit(main())
