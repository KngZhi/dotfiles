"""Patch only XLSX styles and note parts; preserve original values and formula caches."""
from copy import copy
import hashlib
from io import BytesIO
import json
from pathlib import Path
import posixpath
import re
import tempfile
import xml.etree.ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL = 'http://schemas.openxmlformats.org/package/2006/relationships'
DOCREL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
CT = 'http://schemas.openxmlformats.org/package/2006/content-types'
MARK = '\n[container-intake 自动核对]\n'


def annotate(path, issues, expected_hash):
    path = Path(path)
    original = path.read_bytes()
    if hashlib.sha256(original).hexdigest() != expected_hash:
        raise ValueError('标注前文件已变化，重试验证')
    state_path = path.with_suffix('.annotations.json')
    state = json.loads(state_path.read_text()) if state_path.exists() else {}
    wb = openpyxl.load_workbook(BytesIO(original))
    grouped = {}
    for i in issues:
        if i['sheet'] in wb and re.fullmatch(r'[A-Z]+[1-9][0-9]*', i['cell']):
            grouped.setdefault((i['sheet'], i['cell']), []).append(i['message'])
    # Restore only our own unchanged marking, preserving user notes and edits.
    for key, record in state.items():
        sheet, address = key.split('!', 1)
        if sheet not in wb:
            continue
        cell = wb[sheet][address]
        if cell.comment and MARK in cell.comment.text:
            if cell.style_id == record['marked_style']:
                cell._style = copy(wb._cell_styles[record['original_style']])
            user_text = cell.comment.text.split(MARK, 1)[0]
            cell.comment = Comment(user_text, record['author'] or cell.comment.author) if user_text else None
    updated = {}
    for (sheet, address), messages in grouped.items():
        cell = wb[sheet][address]
        old = cell.comment
        record = dict(original_style=cell.style_id, comment=old.text if old else '', author=old.author if old else '')
        cell._style = copy(cell._style)
        cell.fill = PatternFill('solid', fgColor='FFFFC7CE')
        font = copy(cell.font)
        font.color = 'FF9C0006'
        cell.font = font
        note = '\n'.join(dict.fromkeys(messages)) + '\n处理：补齐或对照原件核实后重新验证；未解决前不得视为可导入。'
        cell.comment = Comment(record['comment'] + MARK + note, record['author'] or 'container-intake')
        record['marked_style'] = cell.style_id
        updated[f'{sheet}!{address}'] = record
    if not grouped and not state:
        return 0
    generated = BytesIO()
    wb.save(generated)
    with ZipFile(BytesIO(original)) as src, ZipFile(generated) as new:
        parts = {n: src.read(n) for n in src.namelist()}
        parts['xl/styles.xml'] = new.read('xl/styles.xml')
        old_book = ET.fromstring(parts['xl/workbook.xml'])
        old_rels = ET.fromstring(parts['xl/_rels/workbook.xml.rels'])
        targets = {r.get('Id'): r.get('Target') for r in old_rels}
        for index, sheet in enumerate(old_book.find(f'{{{NS}}}sheets'), 1):
            target = targets[sheet.get(f'{{{DOCREL}}}id')]
            name = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/' + target)
            fresh = f'xl/worksheets/sheet{index}.xml'
            old_xml, new_xml = ET.fromstring(parts[name]), ET.fromstring(new.read(fresh))
            styles = {c.get('r'): c.get('s') for c in new_xml.iter(f'{{{NS}}}c')}
            cells = {c.get('r'): c for c in old_xml.iter(f'{{{NS}}}c')}
            data = old_xml.find(f'{{{NS}}}sheetData')
            for address, style in styles.items():
                cell = cells.get(address)
                if cell is None:
                    # Missing cells need a note/fill but no fabricated value.
                    if not style:
                        continue
                    rownum = re.search(r'\d+', address).group()
                    row = next((r for r in data if r.get('r') == rownum), None)
                    if row is None:
                        row = ET.SubElement(data, f'{{{NS}}}row', {'r': rownum})
                    cell = ET.SubElement(row, f'{{{NS}}}c', {'r': address})
                if style:
                    cell.set('s', style)
                else:
                    cell.attrib.pop('s', None)
            relname = posixpath.dirname(name) + '/_rels/' + posixpath.basename(name) + '.rels'
            freshrel = f'xl/worksheets/_rels/sheet{index}.xml.rels'
            rels = ET.fromstring(parts[relname]) if relname in parts else ET.Element(f'{{{REL}}}Relationships')
            for rel in list(rels):
                if rel.get('Type', '').endswith(('/comments', '/vmlDrawing')):
                    if rel.get('Type', '').endswith('/vmlDrawing'):
                        v = rel.get('Target')
                        v = v.lstrip('/') if v.startswith('/') else posixpath.normpath(posixpath.dirname(name) + '/' + v)
                        if v in parts:
                            vml = ET.fromstring(parts[v])
                            if any(e.get('ObjectType') not in (None, 'Note') for e in vml.iter()):
                                raise ValueError('含VML控件，保留原文件并要求人工标注')
                    rels.remove(rel)
            for drawing in list(old_xml.findall(f'{{{NS}}}legacyDrawing')):
                old_xml.remove(drawing)
            if freshrel in new.namelist():
                for rel in ET.fromstring(new.read(freshrel)):
                    if not rel.get('Type', '').endswith(('/comments', '/vmlDrawing')):
                        continue
                    rel.set('Id', 'containerIntake' + str(len(rels)))
                    t = rel.get('Target')
                    t = t.lstrip('/') if t.startswith('/') else posixpath.normpath(posixpath.dirname(fresh) + '/' + t)
                    rel.set('Target', '/' + t)
                    parts[t] = new.read(t)
                    rels.append(rel)
                    if rel.get('Type').endswith('/vmlDrawing'):
                        ET.SubElement(old_xml, f'{{{NS}}}legacyDrawing', {f'{{{DOCREL}}}id': rel.get('Id')})
            parts[name] = ET.tostring(old_xml, encoding='utf-8', xml_declaration=True)
            parts[relname] = ET.tostring(rels, encoding='utf-8', xml_declaration=True)
        types = ET.fromstring(parts['[Content_Types].xml'])
        for t in ET.fromstring(new.read('[Content_Types].xml')):
            if 'comments' in t.get('ContentType', '') or t.get('Extension') == 'vml':
                if not any(dict(t.attrib) == dict(x.attrib) for x in types):
                    types.append(t)
        parts['[Content_Types].xml'] = ET.tostring(types, encoding='utf-8', xml_declaration=True)
    with tempfile.NamedTemporaryFile(dir=path.parent, suffix='.xlsx', delete=False) as temp:
        tmp = Path(temp.name)
    try:
        with ZipFile(tmp, 'w', compression=ZIP_DEFLATED) as output:
            for name, content in parts.items():
                output.writestr(name, content)
        if path.read_bytes() != original:
            raise ValueError('标注期间文件变化，未覆盖')
        backup = path.with_suffix('.before-annotations.xlsx')
        if not backup.exists():
            backup.write_bytes(original)
        tmp.replace(path)
        state_path.write_text(json.dumps(updated, ensure_ascii=False, indent=2))
    finally:
        tmp.unlink(missing_ok=True)
    return len(grouped)
