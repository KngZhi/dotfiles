#!/usr/bin/env -S uv run --quiet --script
# /// script
# requires-python = ">=3.11"
# dependencies = ["gspread>=6.2,<7", "openpyxl>=3.1"]
# ///
"""Google Sheets working copy of a container table: create, dump, write, set, validate, mark, unmark, export, template.

Run with `uv run scripts/sheets.py <command> ...` from the skill directory. Google API calls are
confined to the functions taking a spreadsheet or client; layout and note planning are pure.
"""
import argparse
import datetime as dt
import json
import os
from pathlib import Path
import re
import sys

from openpyxl.utils import column_index_from_string

sys.path.insert(0, str(Path(__file__).parent))
from grid import book_from_matrices, load_xlsx  # noqa: E402
from validate_workbook import EXIT_CODES, SCHEMA, STANDARD_HEADERS, strip_trailing, validate, validate_book  # noqa: E402

HEADERS = STANDARD_HEADERS
TEXT_COLUMNS = 2  # 货号, 条形码 stay text so barcodes never become numbers
CONFIG = SCHEMA['config']
CONFIG_ROW = {k: i + 2 for i, k in enumerate(CONFIG)}
DATE_PARAMS = ('发柜日期', 'ETA')
CONFIG_DEFAULTS = {
    '货柜号': (None, ''),
    '发柜日期': (None, '实际发柜日期'),
    'ETA': (None, '预计到港日期'),
    '海运费': (None, 'USD，待补'),
    '内陆费': (None, 'CNY；厂家承担填0'),
    '卸柜费': (125000, 'CLP，默认'),
    '清关杂费': (1500000, 'CLP，暂估'),
    'USD-CLP': (None, 'CLP/USD'),
    'USD-CNY': (None, 'CNY/USD，中行现汇卖出价'),
    'CNY-CLP': ('formula', '含12比索/美元手续费'),
    'IVA': ('formula', 'CLP，估算'),
}
CONTAINER_NO = re.compile(r'[A-Z]{4}\d{7}')
CELL = re.compile(r'[A-Z]{1,2}[1-9]\d*')
MARKS = {'ai': '\n[container-intake AI核对]\n', 'auto': '\n[container-intake 自动核对]\n'}
RED = {'backgroundColor': {'red': 1, 'green': 0.8, 'blue': 0.8},
       'textFormat': {'foregroundColor': {'red': 0.6, 'green': 0, 'blue': 0}}}
PLAIN = {'backgroundColor': {'red': 1, 'green': 1, 'blue': 1},
         'textFormat': {'foregroundColor': {'red': 0, 'green': 0, 'blue': 0}}}
MARK_FIELDS = 'note,userEnteredFormat.backgroundColor,userEnteredFormat.textFormat.foregroundColor'

# ---- layout (pure) ---------------------------------------------------------

END = 'MATCH("DATA_END",A:A,0)-1'


def rng(col):
    """Detail range that follows row insertions: header+1 up to the row before DATA_END."""
    return f'{col}2:INDEX({col}:{col},{END})'


def config_formulas():
    b = {k: f'B{CONFIG_ROW[k]}' for k in CONFIG}
    total = 'INDEX(data!L:L,MATCH("合计",data!A:A,0))'
    return {
        'CNY-CLP': f'=IF(COUNT({b["USD-CLP"]},{b["USD-CNY"]})<>2,"",({b["USD-CLP"]}+12)/{b["USD-CNY"]})',
        'IVA': f'=IF(COUNT({total},{b["海运费"]},{b["USD-CLP"]},{b["USD-CNY"]})<>4,"",'
               f'({total}/{b["USD-CNY"]}+{b["海运费"]})*{b["USD-CLP"]}*0.3*0.19)',
    }


def cell_value(v):
    if v is None:
        return ''
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    return v


def data_matrix(rows):
    out = [list(HEADERS)]
    for i, row in enumerate(rows, start=2):
        line = [cell_value(row.get(h)) for h in HEADERS]
        for c in range(TEXT_COLUMNS):
            if line[c] != '':
                line[c] = str(line[c])
        if line[9] == '':
            line[9] = 0
        line[10] = f'=IF(COUNT(H{i}:J{i})<>3,"",H{i}*I{i}+J{i})'
        line[11] = f'=IF(COUNT(G{i},K{i})<>2,"",G{i}*K{i})'
        out.append(line)
    out.append(['DATA_END'])
    out.append([])
    total = [''] * len(HEADERS)
    total[0], total[8], total[12] = '合计', f'=SUM({rng("I")})', f'=SUM({rng("M")})'
    total[11] = f'=IF(OR(COUNTA({rng("A")})=0,COUNT({rng("L")})<>COUNTA({rng("A")})),"",SUM({rng("L")}))'
    out.append(total)
    for label, unit in (('打数合计', '打'), ('条数合计', '条')):
        line = [''] * len(HEADERS)
        line[0], line[5] = label, unit
        line[9] = f'=SUMIF({rng("F")},"{unit}",{rng("J")})'
        line[10] = f'=SUMIF({rng("F")},"{unit}",{rng("K")})'
        out.append(line)
    return out


def config_matrix(values=None, notes=None):
    values, notes = values or {}, notes or {}
    formulas = config_formulas()
    out = [['参数', '值', '说明']]
    for k in CONFIG:
        default, note = CONFIG_DEFAULTS[k]
        v = values.get(k, formulas[k] if default == 'formula' else default)
        out.append([k, cell_value(v), notes.get(k) or note])
    return out


def rows_from_book(book):
    """Extract detail rows, config values and config notes from a standard 13-column book."""
    ws = book['data']
    if strip_trailing(ws.row_values(1)) != HEADERS:
        raise SystemExit('只支持13列标准表；旧表先按工作簿契约迁移')
    ends = [r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == 'DATA_END']
    if len(ends) != 1:
        raise SystemExit('必须且只能有一个 DATA_END')
    rows = [{h: cell_value(ws.cell(r, c + 1).value) for c, h in enumerate(HEADERS)} for r in range(2, ends[0])]
    config, notes = {}, {}
    cfg = book['config']
    for r in range(2, cfg.max_row + 1):
        key = cfg.cell(r, 1).value
        if isinstance(key, str):
            config[key] = cell_value(cfg.cell(r, 2).value)
            notes[key] = cell_value(cfg.cell(r, 3).value)
    return rows, config, notes


def parse_value(text):
    if text == '':
        return ''
    for cast in (int, float):
        try:
            return cast(text)
        except ValueError:
            pass
    return text


def set_cell_value(sheet, cell, raw):
    """Value to write for `set data!<cell>=<raw>`. 货号/条形码 (columns A, B on data) must
    always stay text: parsing them as int/float and sending a JSON number strips Sheets'
    TEXT number format from the cell and turns the barcode into a number for good."""
    if sheet == 'data' and column_index_from_string(re.match('[A-Z]+', cell).group()) <= TEXT_COLUMNS:
        return raw
    return parse_value(raw)


# ---- notes and marks (pure) ---------------------------------------------------


def parse_note(text):
    positions = sorted((text.find(m), owner) for owner, m in MARKS.items() if m in text)
    if not positions:
        return text, {}
    sections = {}
    for i, (pos, owner) in enumerate(positions):
        start = pos + len(MARKS[owner])
        end = positions[i + 1][0] if i + 1 < len(positions) else len(text)
        sections[owner] = text[start:end]
    return text[:positions[0][0]], sections


def render_note(user, sections):
    return user + ''.join(MARKS[o] + sections[o] for o in MARKS if sections.get(o))


def plan_marks(existing, issues, owner):
    """existing: {(sheet, cell): note}; issues: [{sheet, cell, message, blocking?}].
    Returns {(sheet, cell): (note, red)} for every cell whose `owner` section changes."""
    fresh = {}
    for i in issues:
        fresh.setdefault((i['sheet'], i['cell']), []).append(('阻断导入：' if i.get('blocking') else '') + i['message'])
    plan = {}
    for key in set(existing) | set(fresh):
        user, sections = parse_note(existing.get(key, ''))
        before = sections.get(owner)
        after = '\n'.join(fresh[key]) if key in fresh else None
        if before == after:
            continue
        if after is None:
            sections.pop(owner, None)
        else:
            sections[owner] = after
        plan[key] = (render_note(user, sections), bool(sections))
    return plan


def markable(issues, sheet_titles):
    return [i for i in issues if i['sheet'] in sheet_titles and CELL.fullmatch(i['cell'] or '')]


# ---- Google API ---------------------------------------------------------------


SCOPES = ['https://www.googleapis.com/auth/spreadsheets', 'https://www.googleapis.com/auth/drive.file']


def connect():
    import gspread
    from gspread.auth import DEFAULT_CREDENTIALS_FILENAME
    if not Path(DEFAULT_CREDENTIALS_FILENAME).exists():
        raise SystemExit(f'缺少 {DEFAULT_CREDENTIALS_FILENAME}；按 references/google-sheets.md 完成一次性授权')
    return gspread.oauth(scopes=SCOPES)  # drive.file: only files this app created are visible


def folder_id(gc, create=False):
    name = os.environ.get('CONTAINER_INTAKE_DRIVE_FOLDER', '货柜整理')
    session = gc.http_client.session
    query = f"name = '{name}' and mimeType = 'application/vnd.google-apps.folder' and trashed = false"
    r = session.get('https://www.googleapis.com/drive/v3/files', params={'q': query, 'fields': 'files(id)'})
    r.raise_for_status()
    files = r.json().get('files', [])
    if files:
        return files[0]['id']
    if not create:
        return None
    r = session.post('https://www.googleapis.com/drive/v3/files', json={'name': name, 'mimeType': 'application/vnd.google-apps.folder'})
    r.raise_for_status()
    return r.json()['id']


def open_sheet(gc, ref):
    import gspread
    try:
        if ref.startswith('http'):
            return gc.open_by_url(ref)
        if CONTAINER_NO.fullmatch(ref):
            return gc.open(ref, folder_id=folder_id(gc))
        return gc.open_by_key(ref)
    except gspread.SpreadsheetNotFound:
        raise SystemExit(f'找不到表格 {ref}；用 create 新建或传入完整链接')


def create_sheet(gc, title, rows, config, notes):
    sh = gc.create(title, folder_id=folder_id(gc, create=True))
    data = sh.sheet1
    data.update_title('data')
    data.resize(cols=len(HEADERS))
    cfg = sh.add_worksheet('config', rows=len(CONFIG) + 1, cols=3)
    requests = [
        {'repeatCell': {'range': {'sheetId': data.id, 'startColumnIndex': 0, 'endColumnIndex': TEXT_COLUMNS},
                        'cell': {'userEnteredFormat': {'numberFormat': {'type': 'TEXT'}}},
                        'fields': 'userEnteredFormat.numberFormat'}},
        {'updateSheetProperties': {'properties': {'sheetId': data.id, 'gridProperties': {'frozenRowCount': 1}},
                                   'fields': 'gridProperties.frozenRowCount'}},
    ]
    for key in DATE_PARAMS:
        r = CONFIG_ROW[key] - 1
        requests.append({'repeatCell': {'range': {'sheetId': cfg.id, 'startRowIndex': r, 'endRowIndex': r + 1,
                                                  'startColumnIndex': 1, 'endColumnIndex': 2},
                                        'cell': {'userEnteredFormat': {'numberFormat': {'type': 'DATE', 'pattern': 'yyyy-mm-dd'}}},
                                        'fields': 'userEnteredFormat.numberFormat'}})
    sh.batch_update({'requests': requests})
    data.update(data_matrix(rows), 'A1', raw=False)
    cfg.update(config_matrix(config, notes), 'A1', raw=False)
    return sh


def load_book(sh):
    titles = [ws.title for ws in sh.worksheets()]
    ranges = [f"'{t}'" for t in titles]
    computed = sh.values_batch_get(ranges, params={'valueRenderOption': 'UNFORMATTED_VALUE', 'dateTimeRenderOption': 'FORMATTED_STRING'})
    formulas = sh.values_batch_get(ranges, params={'valueRenderOption': 'FORMULA'})
    values = {t: vr.get('values', []) for t, vr in zip(titles, computed['valueRanges'])}
    forms = {t: vr.get('values', []) for t, vr in zip(titles, formulas['valueRanges'])}
    return book_from_matrices(values, forms)


def existing_notes(sh):
    from gspread.utils import rowcol_to_a1
    meta = sh.fetch_sheet_metadata({'includeGridData': True,
                                    'fields': 'sheets(properties(title),data(startRow,startColumn,rowData(values(note))))'})
    notes = {}
    for s in meta['sheets']:
        title = s['properties']['title']
        for block in s.get('data', []):
            r0, c0 = block.get('startRow', 0), block.get('startColumn', 0)
            for ri, row in enumerate(block.get('rowData', [])):
                for ci, cell in enumerate(row.get('values', [])):
                    if cell.get('note'):
                        notes[(title, rowcol_to_a1(r0 + ri + 1, c0 + ci + 1))] = cell['note']
    return notes


def apply_marks(sh, plan):
    from gspread.utils import a1_to_rowcol
    ids = {ws.title: ws.id for ws in sh.worksheets()}
    requests = []
    for (sheet, coord), (note, red) in plan.items():
        r, c = a1_to_rowcol(coord)
        requests.append({'updateCells': {
            'range': {'sheetId': ids[sheet], 'startRowIndex': r - 1, 'endRowIndex': r, 'startColumnIndex': c - 1, 'endColumnIndex': c},
            'rows': [{'values': [{'note': note, 'userEnteredFormat': RED if red else PLAIN}]}],
            'fields': MARK_FIELDS}})
    if requests:
        sh.batch_update({'requests': requests})
    return len(requests)


def mark_issues(sh, issues, owner):
    titles = {ws.title for ws in sh.worksheets()}
    return apply_marks(sh, plan_marks(existing_notes(sh), markable(issues, titles), owner))


def run_validation(sh, report=None):
    result = validate_book(load_book(sh), sh.url)
    changed = mark_issues(sh, result['issues'], 'auto')
    if report:
        report.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    unmarked = [i for i in result['issues'] if not CELL.fullmatch(i['cell'] or '')]
    print(f"{result['status']}: {result['rows']} rows; {len(result['issues'])} issues; {changed} cells re-marked; {sh.url}")
    for i in unmarked:
        print(f"  [{i['level']}] {i['sheet']}!{i['cell']}: {i['message']}")
    return result


# ---- commands -------------------------------------------------------------------


def cmd_auth(args):
    gc = connect()
    print(f'已授权；工作目录 Drive 文件夹：{os.environ.get("CONTAINER_INTAKE_DRIVE_FOLDER", "货柜整理")}（id {folder_id(gc, create=True)}）')


def cmd_create(args):
    if not CONTAINER_NO.fullmatch(args.container):
        raise SystemExit('柜号须为四个大写字母加七位数字')
    rows, config, notes = ([{}], {}, {})
    if args.source:
        rows, config, notes = rows_from_book(load_xlsx(args.source))
    config.setdefault('货柜号', args.container)
    gc = connect()
    sh = create_sheet(gc, args.container, rows, config, notes)
    print(sh.url)
    run_validation(sh)


def cmd_url(args):
    print(open_sheet(connect(), args.ref).url)


def cmd_dump(args):
    sh = open_sheet(connect(), args.ref)
    rows, config, notes = rows_from_book(load_book(sh))
    payload = dict(url=sh.url, rows=rows, config=config, config_notes=notes,
                   cell_notes={f'{s}!{c}': n for (s, c), n in existing_notes(sh).items()})
    text = json.dumps(payload, ensure_ascii=False, indent=2, default=str)
    if args.out:
        args.out.write_text(text + '\n', encoding='utf-8')
        print(args.out)
    else:
        print(text)


def cmd_write(args):
    payload = json.loads(args.rows.read_text(encoding='utf-8'))
    rows = payload['rows'] if isinstance(payload, dict) else payload
    unknown = {k for r in rows for k in r} - set(HEADERS)
    if unknown:
        raise SystemExit(f'未知列：{sorted(unknown)}；只接受 {HEADERS}')
    sh = open_sheet(connect(), args.ref)
    data = sh.worksheet('data')
    data.clear()
    data.update(data_matrix(rows), 'A1', raw=False)
    print(f'已重写 {len(rows)} 行明细')
    run_validation(sh)


def cmd_set(args):
    sh = open_sheet(connect(), args.ref)
    updates = {}
    config_keys = None
    for assignment in args.assignments:
        target, _, raw = assignment.partition('=')
        if target.startswith('config.'):
            value = parse_value(raw)
            key = target[len('config.'):]
            if config_keys is None:
                config_keys = sh.worksheet('config').col_values(1)
            if key not in config_keys:
                raise SystemExit(f'config 没有参数 {key}')
            updates.setdefault('config', []).append({'range': f'B{config_keys.index(key) + 1}', 'values': [[value]]})
        else:
            sheet, _, cell = target.partition('!')
            if not cell or not CELL.fullmatch(cell):
                raise SystemExit(f'目标须为 sheet!A1 或 config.参数：{target}')
            value = set_cell_value(sheet, cell, raw)
            updates.setdefault(sheet, []).append({'range': cell, 'values': [[value]]})
    for sheet, batch in updates.items():
        sh.worksheet(sheet).batch_update(batch, value_input_option='USER_ENTERED')
    print(f'已写入 {sum(len(b) for b in updates.values())} 个单元格')
    run_validation(sh)


def cmd_validate(args):
    sh = open_sheet(connect(), args.ref)
    return EXIT_CODES[run_validation(sh, args.report)['status']]


def cmd_mark(args):
    sheet, _, cell = args.target.partition('!')
    if not cell or not CELL.fullmatch(cell):
        raise SystemExit('目标须为 sheet!A1')
    sh = open_sheet(connect(), args.ref)
    if sheet not in {ws.title for ws in sh.worksheets()}:
        raise SystemExit(f'没有工作表 {sheet}')
    notes = existing_notes(sh)
    _, sections = parse_note(notes.get((sheet, cell), ''))
    kept = [dict(sheet=sheet, cell=cell, message=m) for m in sections.get('ai', '').split('\n') if m]
    if not args.replace:
        kept = [k for k in kept if k['message'] != args.message]
    else:
        kept = []
    kept.append(dict(sheet=sheet, cell=cell, message=args.message, blocking=args.blocking))
    apply_marks(sh, plan_marks(notes, kept, 'ai'))
    print(f'已标注 {sheet}!{cell}')


def cmd_unmark(args):
    sh = open_sheet(connect(), args.ref)
    notes = existing_notes(sh)
    if args.targets:
        wanted = {tuple(t.split('!', 1)) for t in args.targets}
        keep = [dict(sheet=s, cell=c, message=parse_note(n)[1]['ai'])
                for (s, c), n in notes.items() if 'ai' in parse_note(n)[1] and (s, c) not in wanted]
    else:
        keep = []
    # Rebuild the AI layer from what should remain; plan_marks drops the rest.
    issues = [dict(sheet=k['sheet'], cell=k['cell'], message=m) for k in keep for m in k['message'].split('\n') if m]
    changed = apply_marks(sh, plan_marks(notes, issues, 'ai'))
    print(f'已清除 {changed} 个 AI 标注')


def cmd_export(args):
    from gspread.utils import ExportFormat
    gc = connect()
    sh = open_sheet(gc, args.ref)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_bytes(gc.export(sh.id, ExportFormat.EXCEL))
    result = validate(args.out)
    report = args.out.with_suffix('.validation.json')
    report.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f"已导出 {args.out}；{result['status']}: {result['rows']} rows; {len(result['issues'])} issues; {report}")
    return EXIT_CODES[result['status']]


def cmd_template(args):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'data'
    for row in data_matrix([{}]):
        ws.append([None if v == '' else v for v in row])
    for col in ('A', 'B'):
        for cell in ws[col]:
            cell.number_format = '@'
    cfg = wb.create_sheet('config')
    for row in config_matrix():
        cfg.append([None if v == '' else v for v in row])
    for key in DATE_PARAMS:
        cfg.cell(CONFIG_ROW[key], 2).number_format = 'yyyy-mm-dd'
    wb.save(args.out)
    print(args.out)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest='command', required=True)
    p = sub.add_parser('auth', help='完成或刷新 Google 授权并确认工作文件夹')
    p.set_defaults(fn=cmd_auth)
    p = sub.add_parser('create', help='按柜号新建标准表格（可从13列 xlsx 导入）')
    p.add_argument('container')
    p.add_argument('--from', dest='source', type=Path)
    p.set_defaults(fn=cmd_create)
    p = sub.add_parser('url', help='打印表格链接')
    p.add_argument('ref')
    p.set_defaults(fn=cmd_url)
    p = sub.add_parser('dump', help='以 JSON 读出明细、config 和批注')
    p.add_argument('ref')
    p.add_argument('--out', type=Path)
    p.set_defaults(fn=cmd_dump)
    p = sub.add_parser('write', help='用 JSON 明细重写 data 表并验证')
    p.add_argument('ref')
    p.add_argument('rows', type=Path)
    p.set_defaults(fn=cmd_write)
    p = sub.add_parser('set', help='写入单元格或 config 参数并验证：data!G5=12.5 config.海运费=2000')
    p.add_argument('ref')
    p.add_argument('assignments', nargs='+')
    p.set_defaults(fn=cmd_set)
    p = sub.add_parser('validate', help='运行表内验证并把问题标到单元格')
    p.add_argument('ref')
    p.add_argument('--report', type=Path)
    p.set_defaults(fn=cmd_validate)
    p = sub.add_parser('mark', help='AI 核对后给单元格加红标和批注')
    p.add_argument('ref')
    p.add_argument('target', help='sheet!A1')
    p.add_argument('message')
    p.add_argument('--blocking', action='store_true', help='注明阻断导入')
    p.add_argument('--replace', action='store_true', help='替换该单元格已有的 AI 批注')
    p.set_defaults(fn=cmd_mark)
    p = sub.add_parser('unmark', help='清除 AI 标注（不给单元格则清全部）')
    p.add_argument('ref')
    p.add_argument('targets', nargs='*')
    p.set_defaults(fn=cmd_unmark)
    p = sub.add_parser('export', help='导出 xlsx 并做最终表内验证')
    p.add_argument('ref')
    p.add_argument('out', type=Path)
    p.set_defaults(fn=cmd_export)
    p = sub.add_parser('template', help='生成 assets/模板.xlsx')
    p.add_argument('out', type=Path)
    p.set_defaults(fn=cmd_template)
    args = parser.parse_args(argv)
    return args.fn(args) or 0


if __name__ == '__main__':
    raise SystemExit(main())
